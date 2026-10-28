"""
Ingestion: loaders for every supported format, cleaning, chunking
strategies, metadata extraction, and the vector store round trip.

Test types:
- Loader/chunker/metadata tests: real parsers on files generated in a
  temp directory; no models.
- Pipeline tests: real pipeline and a real Chroma store in a temp
  directory, with a deterministic fake embedding function (no model).
"""

import csv
import json
import zipfile
from pathlib import Path

import pytest

from knowledge_base import chunker, pipeline
from knowledge_base.chunker import (
    STRATEGY_DOCUMENT,
    STRATEGY_SEMANTIC,
    _heading_level,
    chunk_document,
)
from knowledge_base.cleaner import is_usable_text, remove_repeated_headers_footers, unwrap_lines
from knowledge_base.file_loader import discover_files
from knowledge_base.metadata_extractor import extract_explicit_fields, extract_metadata
from knowledge_base.parsers import parse_document
from knowledge_base.vector_store import VectorStore


# --------------------------------------------------------------------
# Fixture files, one per format
# --------------------------------------------------------------------

def make_pdf(path: Path, pages: list[list[str]]) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path), pagesize=A4)
    for lines in pages:
        y = 800
        for line in lines:
            pdf.drawString(50, y, line)
            y -= 16
        pdf.showPage()
    pdf.save()
    return path


def make_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("Course Handbook", level=1)
    doc.add_paragraph("General information for learners.")
    doc.add_heading("Fees", level=2)
    table = doc.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("Course", "Fee"), ("Python Basics", "3000 EGP"), ("SQL", "2500 EGP")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.save(str(path))
    return path


def make_pptx(path: Path) -> Path:
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Welcome Slide"
    slide.placeholders[1].text = "Introduction to the programme structure."
    deck.save(str(path))
    return path


def make_xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Schedules"
    sheet.append(["Course ID", "Course Name", "Start Date", "Fee"])
    sheet.append(["PY-101", "Python Basics", "2026-11-01", 3000])
    sheet.append(["SQL-201", "SQL Analytics", "2026-12-01", 2500])
    sheet.append([None, None, None, None])  # empty row is skipped
    book.save(str(path))
    return path


@pytest.fixture
def docs(tmp_path):
    d = tmp_path / "docs"
    d.mkdir()
    make_pdf(d / "guide.pdf", [
        ["Andalusia Test Catalogue", "1. Python Basics", "Level: Beginner", "Delivery mode: Online",
         "1.1 Entry Requirements", "No prior programming experience is needed for this course."],
        ["Andalusia Test Catalogue", "1.2 Learning Outcomes", "Write small Python programs that read",
         "files and print results.", "2. SQL Analytics", "Level: Intermediate"],
    ])
    make_docx(d / "handbook.docx")
    make_pptx(d / "intro.pptx")
    make_xlsx(d / "schedule.xlsx")
    (d / "notes.md").write_text("# Notes\n\n## Refunds\n\nRefunds are possible within 14 days.\n", encoding="utf-8")
    (d / "plain.txt").write_text("Office hours are Sunday to Thursday, 9 to 5.", encoding="utf-8")
    with (d / "instructors.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["instructor_id", "name", "expertise", "years_experience"])
        w.writerow(["I-1", "Mona Adel", "Data analysis", "8"])
        w.writerow(["I-2", "Karim Nabil", "Web development", "11"])
    (d / "faqs.json").write_text(json.dumps({
        "dataset": "FAQ", "updated": "2026-10-01",
        "faqs": [{"id": "F1", "question": "Is there a certificate?", "answer": "Yes, on completion."},
                 {"id": "F2", "question": "Can I pay in instalments?", "answer": "Yes, in two payments."}],
    }), encoding="utf-8")
    (d / "page.html").write_text(
        "<html><head><title>Admissions</title><script>var x=1;</script></head><body>"
        "<nav>Home | About</nav><h1>Admissions</h1><p>Apply online before the start date.</p>"
        "<h2>Documents</h2><ul><li>National ID copy</li><li>Latest certificate</li></ul>"
        "<table><tr><th>Step</th><th>Detail</th></tr><tr><td>1</td><td>Fill the form</td></tr></table>"
        "</body></html>", encoding="utf-8")
    return d


# --------------------------------------------------------------------
# Loaders
# --------------------------------------------------------------------

def test_every_supported_format_is_discovered(docs):
    (docs / "~$lock.xlsx").write_bytes(b"x")
    (docs / "image.png").write_bytes(b"\x89PNG")

    names = {p.name for p in discover_files(docs)}

    assert names == {"guide.pdf", "handbook.docx", "intro.pptx", "schedule.xlsx", "notes.md",
                     "plain.txt", "instructors.csv", "faqs.json", "page.html"}


def test_pdf_keeps_page_numbers(docs):
    doc = parse_document(docs / "guide.pdf")

    assert [p["page"] for p in doc["pages"]] == [1, 2]
    assert "Python Basics" in doc["pages"][0]["text"]


def test_docx_keeps_headings_and_table_rows_in_order(docs):
    text = parse_document(docs / "handbook.docx")["pages"][0]["text"]

    assert text.index("# Course Handbook") < text.index("## Fees") < text.index("Course: Python Basics")
    assert "Course: Python Basics\nFee: 3000 EGP" in text  # row relationship kept


def test_pptx_one_page_per_slide(docs):
    doc = parse_document(docs / "intro.pptx")

    assert doc["pages"][0]["page"] == 1
    assert "Welcome Slide" in doc["pages"][0]["text"]


def test_csv_one_record_per_row(docs):
    pages = parse_document(docs / "instructors.csv")["pages"]

    assert [p["record_id"] for p in pages] == ["I-1", "I-2"]
    assert pages[0]["section"] == "Mona Adel"
    assert pages[0]["fields"]["expertise"] == "Data analysis"
    assert "Karim" not in pages[0]["text"]  # rows never merged


def test_xlsx_one_record_per_row_with_sheet(docs):
    pages = parse_document(docs / "schedule.xlsx")["pages"]

    assert [p["record_id"] for p in pages] == ["PY-101", "SQL-201"]
    assert pages[0]["fields"]["sheet"] == "Schedules"
    assert pages[0]["fields"]["Fee"] == 3000


def test_json_records_carry_shared_document_fields(docs):
    pages = parse_document(docs / "faqs.json")["pages"]

    assert [p["record_id"] for p in pages] == ["F1", "F2"]
    assert pages[0]["section"] == "Is there a certificate?"
    assert pages[0]["fields"]["updated"] == "2026-10-01"
    assert pages[0]["fields"]["collection"] == "faqs"


def test_html_visible_structure_only(docs):
    doc = parse_document(docs / "page.html")
    text = doc["pages"][0]["text"]

    assert doc["title"] == "Admissions"
    assert "# Admissions" in text and "## Documents" in text
    assert "- National ID copy" in text
    assert "Step: 1\nDetail: Fill the form" in text
    assert "var x" not in text and "Home | About" not in text


def test_markdown_and_text(docs):
    assert "## Refunds" in parse_document(docs / "notes.md")["pages"][0]["text"]
    assert "Office hours" in parse_document(docs / "plain.txt")["pages"][0]["text"]


def test_arabic_text_files_are_read(tmp_path):
    path = tmp_path / "ar.txt"
    path.write_text("الدورة متاحة أونلاين.", encoding="cp1256")

    assert "أونلاين" in parse_document(path)["pages"][0]["text"]


@pytest.mark.parametrize("name, content, error", [
    ("fake.pdf", b"this is not a pdf", "not a valid .pdf"),
    ("fake.xlsx", b"plain text", "not a valid .xlsx"),
    ("broken.json", b"{not json", "not valid JSON"),
])
def test_corrupted_or_renamed_files_fail_clearly(tmp_path, name, content, error):
    path = tmp_path / name
    path.write_bytes(content)

    with pytest.raises(ValueError, match=error):
        parse_document(path)


def test_unsupported_format_is_rejected(tmp_path):
    path = tmp_path / "image.png"
    path.write_bytes(b"\x89PNG")

    with pytest.raises(ValueError, match="Unsupported"):
        parse_document(path)


def test_broken_zip_office_file_fails(tmp_path):
    path = tmp_path / "broken.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("nothing.txt", "x")

    with pytest.raises(Exception):
        parse_document(path)


# --------------------------------------------------------------------
# Cleaning
# --------------------------------------------------------------------

def test_wrapped_pdf_lines_are_rejoined_but_structure_kept():
    text = "Analyze a tabular dataset and explain each\ntransformation.\n- A bullet\n1.2 Next Heading\nLevel: Beginner"

    assert unwrap_lines(text) == (
        "Analyze a tabular dataset and explain each transformation.\n- A bullet\n1.2 Next Heading\nLevel: Beginner"
    )


def test_repeated_headers_with_page_numbers_are_removed():
    pages = [{"page": n, "text": f"ACADEMY NAME\nCatalogue\nSynthetic notice\nPage {n}\n"
                                 f"Body line one {n}.\nBody line two.\nBody three.\nBody four.\nBody five.\nEnd {n}."}
             for n in range(1, 5)]

    cleaned = remove_repeated_headers_footers(pages)

    assert all("Synthetic notice" not in p["text"] and "Page" not in p["text"] for p in cleaned)
    assert all("Body line one" in p["text"] for p in cleaned)


@pytest.mark.parametrize("text, usable", [
    ("", False), ("   \n  ", False), ("�" * 50, False),
    ("A short but real sentence about courses.", True), ("الدورة متاحة أونلاين لمدة ستة أسابيع", True),
])
def test_unusable_extractions_are_detected(text, usable):
    assert is_usable_text(text) is usable


# --------------------------------------------------------------------
# Chunking
# --------------------------------------------------------------------

@pytest.mark.parametrize("line, level", [
    ("# Title", 1), ("### Deep", 3), ("1. Python Basics", 1), ("1.2 Entry Requirements", 2),
    ("CATALOGUE OVERVIEW", None), ("Entry Requirements", None), ("Learning outcomes", None),
    ("50 guided learning hours", False), ("Level: Beginner", False),
    ("Analyze a tabular dataset and explain each", False),
    ("1. Write a campaign objective connected to a defined need.", False),
])
def test_heading_levels(line, level):
    assert _heading_level(line) == level


def _doc(text, **extra):
    return {"source": "doc.md", "file_type": ".md", "pages": [{"page": 1, "text": text}], **extra}


COURSES = (
    "# Catalogue\n\n## 1. Python Basics\n\nLevel: Beginner\nDelivery mode: Online\n\n"
    "### Entry Requirements\n\nNo prior experience.\n\n### Fees\n\nFee: 3000 EGP\n\n"
    "## 2. SQL Analytics\n\nLevel: Intermediate\n\n### Entry Requirements\n\nBasic spreadsheets.\n"
)


def test_chunks_carry_their_section_path_and_inherited_fields():
    chunks = chunk_document(_doc(COURSES.replace("## 1.", "## ").replace("## 2.", "## ")))

    sql = [c for c in chunks if "Basic spreadsheets" in c["text"]][0]
    assert sql["text"].startswith("Catalogue > SQL Analytics")
    assert sql["fields"]["level"] == "Intermediate"
    python = [c for c in chunks if "No prior experience" in c["text"]][0]
    assert "Python Basics" in python["section_path"]
    assert python["fields"]["delivery_mode"] == "Online"
    assert "Basic spreadsheets" not in python["text"]  # different courses never packed together


def test_small_sections_are_packed_not_left_as_fragments():
    chunks = chunk_document(_doc(COURSES))

    python = [c for c in chunks if "Python Basics" in (c["section_path"] or "")]
    assert len(python) == 1
    assert "Entry Requirements" in python[0]["text"] and "Fee: 3000 EGP" in python[0]["text"]


def test_records_are_never_merged():
    pages = [{"page": None, "record_id": f"R{i}", "section": f"Item {i}", "fields": {"id": f"R{i}"},
              "text": f"Item {i}\nid: R{i}"} for i in range(3)]

    chunks = chunk_document({"source": "x.csv", "file_type": ".csv", "pages": pages})

    assert [c["record_id"] for c in chunks] == ["R0", "R1", "R2"]
    assert all(c["chunk_strategy"] == "record" for c in chunks)


def test_oversized_section_is_split_recursively_with_overlap_and_header():
    body = " ".join(f"Sentence number {i} explains a detail of the topic." for i in range(60))
    chunks = chunk_document(_doc(f"# Guide\n\n## Details\n\n{body}"), max_chars=400, strategy=STRATEGY_DOCUMENT)

    assert len(chunks) > 3
    assert all(len(c["text"]) <= 400 for c in chunks)
    assert all(c["text"].startswith("Guide > Details") for c in chunks)
    assert {c["chunk_strategy"] for c in chunks} == {"recursive"}


def test_semantic_strategy_splits_at_topic_shifts(monkeypatch):
    # Two topics; fake embeddings put each topic's sentences together.
    football = [f"Football match {i} had many goals and fans." for i in range(4)]
    cooking = [f"Cooking recipe {i} uses garlic and olive oil." for i in range(4)]
    monkeypatch.setattr(chunker, "embed_texts",
                        lambda texts: [[1.0, 0.0] if "Football" in t else [0.0, 1.0] for t in texts])

    chunks = chunk_document(_doc("# Notes\n\n" + " ".join(football + cooking)),
                            max_chars=1000, min_chars=50, strategy=STRATEGY_SEMANTIC)

    assert len(chunks) == 2
    assert "Football" in chunks[0]["text"] and "Cooking" not in chunks[0]["text"]
    assert "Cooking" in chunks[1]["text"]
    assert {c["chunk_strategy"] for c in chunks} == {"semantic"}


def test_semantic_strategy_falls_back_when_embeddings_fail(monkeypatch):
    def boom(texts):
        raise RuntimeError("model unavailable")
    monkeypatch.setattr(chunker, "embed_texts", boom)

    chunks = chunk_document(_doc("# N\n\n" + "One sentence here. " * 10), strategy=STRATEGY_SEMANTIC)

    assert chunks and all(c["text"] for c in chunks)


def test_document_strategy_never_calls_the_model(monkeypatch):
    monkeypatch.setattr(chunker, "embed_texts", lambda texts: pytest.fail("model called"))
    body = " ".join(f"Sentence {i} is prose." for i in range(200))

    chunk_document(_doc(f"# Long\n\n{body}"), strategy=STRATEGY_DOCUMENT)


def test_invalid_strategy_is_rejected():
    with pytest.raises(ValueError):
        chunk_document(_doc("text"), strategy="magic")


# --------------------------------------------------------------------
# Metadata
# --------------------------------------------------------------------

def test_explicit_fields_are_discovered_and_typed():
    fields = extract_explicit_fields("Level: Beginner\nFee: 3000\nDuration: 6 weeks\n"
                                     "This line has far too many words before the colon: so it is prose\n"
                                     "المستوى: مبتدئ")

    assert fields == {"level": "Beginner", "fee": 3000, "duration": "6 weeks", "المستوى": "مبتدئ"}


def test_metadata_includes_provenance_and_record_fields():
    meta = extract_metadata({
        "text": "Python Basics\nlevel: Beginner", "source": "knowledge_base/documents/x.csv",
        "file_type": ".csv", "page": None, "section": "Python Basics", "record_id": "PY-101",
        "chunk_id": 3, "fields": {"Course ID": "PY-101", "Fee": "3000"}, "section_path": "Python Basics",
        "chunk_strategy": "record",
    })

    assert meta["filename"] == "x.csv" and meta["record_id"] == "PY-101"
    assert meta["explicit_fields"]["course_id"] == "PY-101"
    assert meta["explicit_fields"]["fee"] == 3000
    assert meta["explicit_fields"]["level"] == "Beginner"


def test_no_metadata_is_invented():
    meta = extract_metadata({"text": "Just a sentence without fields.", "source": "a.txt", "chunk_id": 1})

    assert meta["explicit_fields"] == {}
    assert meta["record_id"] is None and meta["page"] is None


# --------------------------------------------------------------------
# Pipeline + Chroma (fake embeddings, temp store)
# --------------------------------------------------------------------

def fake_embed(texts):
    # Deterministic 8-dimensional vectors from character counts.
    return [[float(t.count(ch) + 1) for ch in "aeiorstn"] for t in texts]


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setattr(pipeline, "embed_texts", fake_embed)
    return VectorStore(persist_directory=tmp_path / "chroma")


def ingest(docs, store, tmp_path):
    return pipeline.ingest_directory(docs, vector_store=store, metadata_path=tmp_path / "meta.json",
                                     manifest_path=tmp_path / "manifest.json")


def test_every_format_is_ingested_end_to_end(docs, store, tmp_path):
    summary = ingest(docs, store, tmp_path)

    assert summary["files_failed"] == []
    assert summary["files_processed"] == 9
    stored = store.collection.get(include=["metadatas"])["metadatas"]
    assert {m["file_type"] for m in stored} == {".pdf", ".docx", ".pptx", ".xlsx", ".md", ".txt",
                                                ".csv", ".json", ".html"}


def test_chroma_metadata_is_flat_filterable_and_portable(docs, store, tmp_path):
    ingest(docs, store, tmp_path)

    rows = store.collection.get(where={"field_expertise": "Data analysis"}, include=["metadatas"])
    assert len(rows["ids"]) == 1
    meta = rows["metadatas"][0]
    assert meta["record_id"] == "I-1"
    assert meta["filename"] == "instructors.csv"
    assert meta["embedding_model"] == pipeline.DEFAULT_MODEL_NAME
    assert all(not isinstance(v, (list, dict)) for v in meta.values())

    pdf = store.collection.get(where={"field_level": "Beginner"}, include=["metadatas"])
    assert pdf["ids"]


def test_reingestion_creates_no_duplicates(docs, store, tmp_path):
    ingest(docs, store, tmp_path)
    first = store.collection.count()

    ingest(docs, store, tmp_path)

    assert store.collection.count() == first


def test_removed_and_shrunk_documents_leave_no_stale_vectors(docs, store, tmp_path):
    ingest(docs, store, tmp_path)
    (docs / "plain.txt").unlink()
    (docs / "instructors.csv").write_text("instructor_id,name\nI-1,Mona Adel\n", encoding="utf-8")

    ingest(docs, store, tmp_path)

    sources = [m["source"] for m in store.collection.get(include=["metadatas"])["metadatas"]]
    assert not any(s.endswith("plain.txt") for s in sources)
    assert sum(s.endswith("instructors.csv") for s in sources) == 1


def test_bad_files_are_reported_and_others_still_indexed(docs, store, tmp_path):
    (docs / "scan.pdf").write_bytes(make_pdf(docs / "scan.pdf", [[]]).read_bytes())  # no text layer
    (docs / "broken.json").write_text("{nope", encoding="utf-8")

    summary = ingest(docs, store, tmp_path)

    failed = {Path(f["source"]).name for f in summary["files_failed"]}
    assert failed == {"scan.pdf", "broken.json"}
    assert summary["files_processed"] == 9


def test_other_collections_in_the_store_are_untouched(docs, store, tmp_path):
    store.add_chunks([{"text": "Approved research page.", "chunk_id": 1, "source": "research_x.md"}],
                     [fake_embed(["Approved research page."])[0]])

    ingest(docs, store, tmp_path)
    ingest(docs, store, tmp_path)

    assert store.collection.get(where={"source": "research_x.md"})["ids"]


def test_sources_inside_the_project_are_stored_relative():
    inside = pipeline.PROJECT_ROOT / "knowledge_base" / "documents" / "x.pdf"

    assert pipeline._relative_source(inside) == str(Path("knowledge_base/documents/x.pdf"))
    assert pipeline._normalize_source("knowledge_base/documents/x.pdf") == str(inside.resolve())
