
"""
Format-specific loaders. Every parser returns the same structure:

    {"source": str, "file_type": ".ext", "pages": [page, ...], ...}

A page is {"page": int | None, "text": str} -- a PDF page, a slide, or
the whole text of a flat document. Structured formats (CSV, XLSX, JSON)
produce one page per record instead, so unrelated rows are never merged:

    {"page": None, "text": "field: value\\n...", "record_id": str,
     "section": record title or None, "fields": {field: value}}

Headings found in DOCX and HTML are written as Markdown headings ("#",
"##") so the chunker sees the document's structure. No parser assumes a
particular dataset: field names come from the file itself.
"""

import csv
import io
import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from knowledge_base.file_loader import check_signature, get_file_type

# Generic field-name conventions (not dataset fields): which column
# identifies a record and which one names it.
_ID_FIELD = re.compile(r"(?:^|[\s_\-.])(?:id|code|key|ref|sku|number|no)$", re.IGNORECASE)
_TITLE_FIELD = re.compile(
    r"(?:^|[\s_\-.])(?:name|title|question|subject|heading|label)$", re.IGNORECASE
)

MAX_FIELD_CHARS = 2000


def parse_document(file_path: str | Path) -> dict:
    """Extract text and relevant content from a supported document."""
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File does not exist: {path}")

    if not path.is_file():
        raise ValueError(f"Path is not a file: {path}")

    extension = get_file_type(path)
    check_signature(path)

    parsers = {
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".pptx": _parse_pptx,
        ".txt": _parse_text,
        ".md": _parse_text,
        ".csv": _parse_csv,
        ".xlsx": _parse_xlsx,
        ".json": _parse_json,
        ".html": _parse_html,
        ".htm": _parse_html,
    }

    return parsers[extension](path)


# --------------------------------------------------------------------
# Shared helpers
# --------------------------------------------------------------------

def _read_text(path: Path) -> str:
    """UTF-8 (with or without BOM), falling back to Windows Arabic."""
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1256"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"{path.name} is not readable text")


def _scalar(value: Any) -> Any:
    """A value as stored in record fields: str/int/float/bool, or None."""
    if value is None:
        return None
    if isinstance(value, bool) or isinstance(value, (int, float)):
        return value
    if isinstance(value, (list, tuple)):
        parts = [str(_scalar(v)) for v in value if _scalar(v) not in (None, "")]
        return "; ".join(parts) or None
    if hasattr(value, "isoformat"):  # dates and datetimes from XLSX
        return value.isoformat()
    text = re.sub(r"[ \t]+", " ", str(value)).strip()
    return text[:MAX_FIELD_CHARS] or None


def _unique_headers(headers: list[Any]) -> list[str]:
    names, seen = [], {}
    for index, header in enumerate(headers, start=1):
        name = str(header).strip() if header is not None else ""
        name = name or f"column_{index}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        names.append(name)
    return names


def _record_page(fields: dict[str, Any], fallback_id: str, context: str = "") -> dict | None:
    """One record as a page: identity first, then every non-empty field."""
    fields = {k: v for k, v in ((k, _scalar(v)) for k, v in fields.items()) if v not in (None, "")}
    if not fields:
        return None

    record_id = next(
        (str(v) for k, v in fields.items() if _ID_FIELD.search(k) or k.lower() == "id"),
        fallback_id,
    )
    title = next((str(v) for k, v in fields.items() if _TITLE_FIELD.search(k)), None)

    lines = [title] if title else []
    lines += [f"{key}: {value}" for key, value in fields.items()]

    return {
        "page": None,
        "record_id": record_id,
        "section": title,
        "fields": fields,
        "text": (f"{context}\n" if context else "") + "\n".join(lines),
    }


def _table_rows_text(rows: list[list[str]]) -> str:
    """A table as one "header: value" block per row (keeps row relationships)."""
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if len(rows) < 2:
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)

    headers = _unique_headers(rows[0])
    blocks = []
    for row in rows[1:]:
        pairs = [f"{h}: {c.strip()}" for h, c in zip(headers, row) if c.strip()]
        if pairs:
            blocks.append("\n".join(pairs))
    return "\n\n".join(blocks)


# --------------------------------------------------------------------
# Page-based formats
# --------------------------------------------------------------------

def _parse_pdf(path: Path) -> dict:
    """Extract text from each PDF page."""
    reader = PdfReader(str(path))
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        pages.append({
            "page": page_number,
            "text": page.extract_text() or "",
        })

    return {
        "source": str(path),
        "file_type": ".pdf",
        "pages": pages,
    }


def _parse_docx(path: Path) -> dict:
    """
    Extract paragraphs and tables in document order. Heading styles
    become Markdown headings; each table row keeps its column headers.
    """
    document = Document(str(path))
    parts: list[str] = []
    tables: list[list[list[str]]] = []

    paragraphs = {id(p._element): p for p in document.paragraphs}
    table_objects = {id(t._element): t for t in document.tables}

    for element in document.element.body.iterchildren():
        if id(element) in paragraphs:
            paragraph = paragraphs[id(element)]
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            level = re.match(r"Heading (\d)", style)
            if style == "Title":
                parts.append(f"# {text}")
            elif level:
                parts.append(f"{'#' * min(int(level.group(1)), 6)} {text}")
            elif "List" in style:
                parts.append(f"- {text}")
            else:
                parts.append(text)
        elif id(element) in table_objects:
            rows = [[cell.text.strip() for cell in row.cells] for row in table_objects[id(element)].rows]
            tables.append(rows)
            rendered = _table_rows_text(rows)
            if rendered:
                parts.append(rendered)

    return {
        "source": str(path),
        "file_type": ".docx",
        "pages": [{
            "page": None,
            "text": "\n\n".join(parts),
        }],
        "tables": tables,
    }


def _parse_pptx(path: Path) -> dict:
    """Extract text and table content from each PowerPoint slide."""
    presentation = Presentation(str(path))
    slides = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        elements = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()

                if text:
                    elements.append(text)

            if shape.has_table:
                for row in shape.table.rows:
                    elements.append(
                        " | ".join(
                            cell.text.strip()
                            for cell in row.cells
                        )
                    )

        slides.append({
            "page": slide_number,
            "text": "\n".join(elements),
        })

    return {
        "source": str(path),
        "file_type": ".pptx",
        "pages": slides,
    }


def _parse_text(path: Path) -> dict:
    """Read a plain-text or Markdown document."""
    text = _read_text(path)

    return {
        "source": str(path),
        "file_type": path.suffix.lower(),
        "pages": [{
            "page": None,
            "text": text,
        }],
    }


def _parse_html(path: Path) -> dict:
    """
    Visible text of an HTML page in reading order: headings as Markdown
    headings, list items as bullets, tables row by row. Scripts, styles
    and navigation are dropped.
    """
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_read_text(path), "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "template", "svg"]):
        tag.decompose()

    title = soup.title.get_text(" ", strip=True) if soup.title else None
    blocks = ("h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table", "pre", "blockquote", "dt", "dd")
    parts: list[str] = []

    for element in (soup.body or soup).find_all(blocks):
        if element.find_parent(["table", "li", "pre"]) is not None and element.name != "table":
            continue  # already rendered by its container
        if element.name == "table" and element.find_parent("table") is not None:
            continue
        if element.name == "table":
            rows = [
                [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
                for row in element.find_all("tr")
            ]
            rendered = _table_rows_text(rows)
        else:
            text = element.get_text(" ", strip=True)
            if not text:
                continue
            if element.name[0] == "h":
                rendered = f"{'#' * int(element.name[1])} {text}"
            elif element.name == "li":
                rendered = f"- {text}"
            else:
                rendered = text
        if rendered:
            parts.append(rendered)

    return {
        "source": str(path),
        "file_type": path.suffix.lower(),
        "title": title,
        "pages": [{
            "page": None,
            "text": "\n\n".join(parts),
        }],
    }


# --------------------------------------------------------------------
# Record-based formats
# --------------------------------------------------------------------

def _parse_csv(path: Path) -> dict:
    """One record per row, keyed by the header row."""
    reader = csv.reader(io.StringIO(_read_text(path)))
    rows = [row for row in reader if any(cell.strip() for cell in row)]

    pages = []
    if rows:
        headers = _unique_headers(rows[0])
        for number, row in enumerate(rows[1:], start=2):
            record = _record_page(dict(zip(headers, row)), f"row {number}")
            if record:
                pages.append(record)

    return {
        "source": str(path),
        "file_type": ".csv",
        "pages": pages,
    }


def _parse_xlsx(path: Path) -> dict:
    """One record per row of every sheet, keyed by that sheet's header row."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    pages = []
    try:
        for sheet in workbook.worksheets:
            rows = [
                list(row) for row in sheet.iter_rows(values_only=True)
                if any(cell not in (None, "") for cell in row)
            ]
            if len(rows) < 2:
                continue
            headers = _unique_headers(rows[0])
            for number, row in enumerate(rows[1:], start=2):
                fields = {"sheet": sheet.title, **dict(zip(headers, row))}
                record = _record_page(fields, f"{sheet.title}!{number}")
                if record:
                    pages.append(record)
    finally:
        workbook.close()

    return {
        "source": str(path),
        "file_type": ".xlsx",
        "pages": pages,
    }


def _flatten(value: Any, prefix: str = "") -> dict[str, Any]:
    """Nested objects as dotted keys; lists of scalars joined."""
    if isinstance(value, dict):
        flat: dict[str, Any] = {}
        for key, item in value.items():
            flat.update(_flatten(item, f"{prefix}.{key}" if prefix else str(key)))
        return flat
    if isinstance(value, list) and any(isinstance(item, (dict, list)) for item in value):
        flat = {}
        for index, item in enumerate(value, start=1):
            flat.update(_flatten(item, f"{prefix}[{index}]"))
        return flat
    return {prefix or "value": value}


def _parse_json(path: Path) -> dict:
    """
    A list of objects -> one record per object. An object holding lists
    of objects -> one record per item, each tagged with its list name and
    the object's top-level values (e.g. dataset name, update date).
    Anything else -> its flattened fields as one record.
    """
    try:
        data = json.loads(_read_text(path))
    except json.JSONDecodeError as exc:
        raise ValueError(f"{path.name} is not valid JSON: {exc}") from exc

    pages: list[dict] = []

    def add_records(items: list, collection: str | None, shared: dict[str, Any]) -> None:
        for number, item in enumerate(items, start=1):
            fields = {**shared}
            if collection:
                fields["collection"] = collection
            fields.update(_flatten(item) if isinstance(item, dict) else {"value": item})
            record = _record_page(fields, f"{collection or 'item'} {number}")
            if record:
                pages.append(record)

    if isinstance(data, list):
        add_records(data, None, {})
    elif isinstance(data, dict):
        lists = {k: v for k, v in data.items() if isinstance(v, list) and v and all(isinstance(i, dict) for i in v)}
        shared = {k: v for k, v in data.items() if not isinstance(v, (list, dict))}
        if lists:
            for name, items in lists.items():
                add_records(items, name, shared)
        else:
            record = _record_page(_flatten(data), "document")
            if record:
                pages.append(record)

    return {
        "source": str(path),
        "file_type": ".json",
        "pages": pages,
    }
