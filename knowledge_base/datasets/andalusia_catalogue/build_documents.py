"""
Build the Andalusia Academy demonstration catalogue from one structured
source (catalogue_source.json), so every generated document agrees:

    catalogue_<department>.pdf   one per department, one course per section
    instructors.csv              one instructor per row
    schedules_and_fees.xlsx      sheets: Fees, Intakes, Promotions
    faqs.json                    one FAQ per record
    admission_requirements.md    admission policy and process
    course_list.md               one line per course (overview)

All content is SYNTHETIC demonstration data and is labelled as such in
every file. This dataset-specific code is kept apart from the generic
RAG engine in knowledge_base/; the engine ingests the output like any
other documents.

Usage:
    python -m knowledge_base.datasets.andalusia_catalogue.build_documents [--out DIR]
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "catalogue_source.json"
DEFAULT_OUT = HERE.parents[1] / "documents"

ACCREDITATION = "Not accredited (certificate of completion only)"


def load_source(path: Path = SOURCE) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _course_names(data: dict) -> dict[str, str]:
    return {course["id"]: course["name"] for course in data["courses"]}


# --------------------------------------------------------------------
# Course text (shared by the PDFs and the tests)
# --------------------------------------------------------------------

def course_lines(data: dict, course: dict, number: int) -> list[tuple[str, str]]:
    """
    The course section as (style, text) lines. Facts are "Key: value"
    lines and sections are numbered headings ("1.", "1.1"), which survive
    PDF text extraction and give the chunker its structure.
    """
    departments = {d["id"]: d for d in data["departments"]}
    instructors = {i["id"]: i for i in data["instructors"]}
    names = _course_names(data)
    department = departments[course["department"]]
    prereqs = [f"{pid} {names[pid]}" for pid in course["prerequisites"]] or ["None"]

    lines = [
        ("h1", f"{number}. {course['name']}"),
        ("kv", f"Course ID: {course['id']}"),
        ("kv", f"Department: {department['name']}"),
        ("kv", f"Industry: {department['industry']}"),
        ("kv", f"Level: {course['level']}"),
        ("kv", f"Delivery mode: {course['delivery_mode']}"),
        ("kv", f"Language of instruction: {course['language']}"),
        ("kv", f"Duration: {course['duration_weeks']} weeks ({course['total_hours']} hours)"),
        ("kv", f"Time commitment: {course['weekly_commitment']}"),
        ("kv", f"Fee (EGP): {course['fee_egp']}"),
        ("kv", f"Payment options: {course['installments']}"),
        ("kv", f"Certificate: {course['certificate']}"),
        ("kv", f"Accreditation: {ACCREDITATION}"),
        ("kv", f"Country: {data['country']}"),
        ("kv", "Data status: Synthetic demonstration data"),
        ("kv", f"Updated: {data['updated']}"),
        ("h2", f"{number}.1 Course Description"),
        ("p", course["description"]),
        ("kv", f"Target audience: {course['target_audience']}"),
        ("h2", f"{number}.2 Entry Requirements and Prerequisites"),
        ("kv", f"Entry requirements: {course['entry_requirements']}"),
        ("kv", f"Prerequisite courses: {'; '.join(prereqs)}"),
        ("h2", f"{number}.3 Learning Outcomes"),
        *[("li", outcome) for outcome in course["learning_outcomes"]],
        ("h2", f"{number}.4 Topics Covered"),
        *[("li", topic) for topic in course["topics"]],
        ("h2", f"{number}.5 Skills Taught"),
        *[("li", skill) for skill in course["skills"]],
        ("h2", f"{number}.6 Career Outcomes"),
        *[("li", outcome) for outcome in course["career_outcomes"]],
        ("h2", f"{number}.7 Schedule and Intakes"),
        *[("li", f"Starts {i['start']}: {i['schedule']}; {i['seats']} seats") for i in course["intakes"]],
        ("h2", f"{number}.8 Instructors"),
        *[
            ("li", f"{instructors[iid]['name']}, {instructors[iid]['title']} "
                   f"({instructors[iid]['years_experience']} years of experience): {instructors[iid]['expertise']}")
            for iid in course["instructors"]
        ],
    ]
    return lines


# --------------------------------------------------------------------
# PDFs
# --------------------------------------------------------------------

def build_pdfs(data: dict, out: Path) -> list[Path]:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    styles = getSampleStyleSheet()
    style = {
        "title": ParagraphStyle("t", parent=styles["Title"], fontSize=18),
        "h1": ParagraphStyle("h1", parent=styles["Heading1"], fontSize=15, spaceBefore=4),
        "h2": ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, spaceBefore=6),
        "p": ParagraphStyle("p", parent=styles["BodyText"], fontSize=10),
        "kv": ParagraphStyle("kv", parent=styles["BodyText"], fontSize=10, spaceAfter=0, spaceBefore=0),
        "li": ParagraphStyle("li", parent=styles["BodyText"], fontSize=10, leftIndent=12, spaceAfter=0),
        "note": ParagraphStyle("n", parent=styles["BodyText"], fontSize=10, textColor="#8a1c1c"),
    }

    paths = []
    for department in data["departments"]:
        courses = [c for c in data["courses"] if c["department"] == department["id"]]
        if not courses:
            continue
        path = out / f"catalogue_{department['id'].lower()}_{department['name'].lower().replace(' ', '_')}.pdf"
        title = f"Andalusia Academy Course Catalogue: {department['name']}"

        def page_frame(canvas, _doc, dept=department["name"]):
            canvas.saveState()
            canvas.setFont("Helvetica", 8)
            canvas.drawString(2 * cm, A4[1] - 1.2 * cm, f"Andalusia Academy - Course Catalogue - {dept}")
            canvas.drawString(2 * cm, 1.0 * cm, f"Synthetic demonstration data - not real academy information - Page {canvas.getPageNumber()}")
            canvas.restoreState()

        story = [
            Paragraph(escape(title), style["title"]),
            Paragraph(escape(data["data_status"]), style["note"]),
            Spacer(1, 8),
            Paragraph(escape(
                f"This catalogue lists the {department['name']} courses of Andalusia Academy "
                f"({data['academy']['campus']}). Fees are in {data['currency']}. "
                f"Online courses use: {data['academy']['online_platform']}."
            ), style["p"]),
            Spacer(1, 6),
            Paragraph("Courses in this catalogue", style["h2"]),
            *[Paragraph(escape(f"- {c['id']} {c['name']} ({c['level']}, {c['delivery_mode']})"), style["li"]) for c in courses],
        ]
        for number, course in enumerate(courses, start=1):
            story.append(PageBreak())
            for kind, text in course_lines(data, course, number):
                story.append(Paragraph(escape(f"- {text}" if kind == "li" else text), style[kind]))

        SimpleDocTemplate(str(path), pagesize=A4, title=title, author="Andalusia Academy (synthetic demo)",
                          topMargin=2 * cm, bottomMargin=2 * cm).build(
            story, onFirstPage=page_frame, onLaterPages=page_frame)
        paths.append(path)
    return paths


# --------------------------------------------------------------------
# Structured files
# --------------------------------------------------------------------

def build_instructors_csv(data: dict, out: Path) -> Path:
    names = _course_names(data)
    path = out / "instructors.csv"
    with path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["instructor_id", "name", "title", "expertise", "years_experience",
                         "languages", "courses_taught", "bio", "data_status"])
        for instructor in data["instructors"]:
            taught = [f"{c['id']} {names[c['id']]}" for c in data["courses"] if instructor["id"] in c["instructors"]]
            writer.writerow([instructor["id"], instructor["name"], instructor["title"], instructor["expertise"],
                             instructor["years_experience"], instructor["languages"], "; ".join(taught),
                             instructor["bio"], "Synthetic demonstration data"])
    return path


def build_schedules_xlsx(data: dict, out: Path) -> Path:
    from openpyxl import Workbook

    departments = {d["id"]: d["name"] for d in data["departments"]}
    path = out / "schedules_and_fees.xlsx"
    book = Workbook()

    fees = book.active
    fees.title = "Fees"
    fees.append(["course_id", "course_name", "department", "level", "delivery_mode", "fee_egp",
                 "payment_options", "currency", "data_status"])
    for c in data["courses"]:
        fees.append([c["id"], c["name"], departments[c["department"]], c["level"], c["delivery_mode"],
                     c["fee_egp"], c["installments"], data["currency"], "Synthetic demonstration data"])

    intakes = book.create_sheet("Intakes")
    intakes.append(["intake_id", "course_name", "course_code", "start_date", "schedule", "delivery_mode",
                    "seats", "data_status"])
    for c in data["courses"]:
        for intake in c["intakes"]:
            intakes.append([f"{c['id']}-{intake['start']}", c["name"], c["id"], intake["start"], intake["schedule"],
                            c["delivery_mode"], intake["seats"], "Synthetic demonstration data"])

    promos = book.create_sheet("Promotions")
    promos.append(["promo_code", "promotion_name", "discount", "conditions", "valid_until", "applies_to",
                   "data_status"])
    for p in data["promotions"]:
        promos.append([p["code"], p["offer"], p["discount"], p["conditions"], p["valid_until"], p["applies_to"],
                       "Synthetic demonstration data"])

    book.save(str(path))
    return path


def build_faqs_json(data: dict, out: Path) -> Path:
    path = out / "faqs.json"
    path.write_text(json.dumps({
        "dataset": "Andalusia Academy frequently asked questions",
        "data_status": "Synthetic demonstration data",
        "updated": data["updated"],
        "faqs": data["faqs"],
    }, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def build_admission_md(data: dict, out: Path) -> Path:
    admission = data["admission"]
    lines = [
        "# Andalusia Academy Admission Requirements",
        "",
        f"Data status: {data['data_status']}",
        f"Updated: {data['updated']}",
        "",
        "## General Requirements",
        "",
        *[f"- {item}" for item in admission["general_requirements"]],
        "",
        "## Required Documents",
        "",
        *[f"- {item}" for item in admission["documents"]],
        "",
        "## Placement Test",
        "",
        admission["placement"],
        "",
        "## English Language Level",
        "",
        admission["english"],
        "",
        "## Admission Process",
        "",
        *[f"{n}. {step}" for n, step in enumerate(admission["process"], start=1)],
        "",
    ]
    path = out / "admission_requirements.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def build_course_list_md(data: dict, out: Path) -> Path:
    """
    One line per course, so "which courses do you offer?" can be
    answered from a passage that lists every course (a single course
    page cannot).
    """
    # Grouped by level: each group is one short passage, so a listing
    # question retrieves whole groups instead of a list cut in half.
    lines = ["# Andalusia Academy Course List"]
    for level in ("Beginner", "Intermediate", "Advanced"):
        courses = [c for c in data["courses"] if c["level"] == level]
        if courses:
            lines += ["", f"## {level} Courses Offered", ""]
            lines += [
                f"- {c['id']} {c['name']}: {c['delivery_mode'].split(' (')[0]}, {c['fee_egp']} {data['currency']}"
                for c in courses
            ]
    lines += [
        "",
        "## About This List",
        "",
        f"{len(data['courses'])} courses in total; each line gives the course, delivery mode and fee.",
        f"Data status: {data['data_status']}",
        f"Updated: {data['updated']}",
        "",
    ]
    path = out / "course_list.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def build_all(out: Path = DEFAULT_OUT, source: Path = SOURCE) -> list[Path]:
    data = load_source(source)
    out.mkdir(parents=True, exist_ok=True)
    return [
        *build_pdfs(data, out),
        build_instructors_csv(data, out),
        build_schedules_xlsx(data, out),
        build_faqs_json(data, out),
        build_admission_md(data, out),
        build_course_list_md(data, out),
    ]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--out", default=str(DEFAULT_OUT))
    args = parser.parse_args(argv)
    for path in build_all(Path(args.out)):
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
